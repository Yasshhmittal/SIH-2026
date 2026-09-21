"""Deliverable renderers — the "real files, not chat replies" requirement.

Documents are built from a *typed spec*, never from model-authored markup. The
model supplies content; this module owns structure. That separation is what
lets the citation verifier block an unsupported claim before it reaches a file
somebody signs.
"""

from __future__ import annotations

import uuid
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ..agent.schemas import Observation
from .base import CAP_RENDER, ToolContext, ToolSpec, registry


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT


def _findings_table(doc: Document, findings: list[dict[str, Any]]) -> None:
    if not findings:
        return
    cols = ["Location", "Observation", "Measured", "Minimum", "Severity"]
    table = doc.add_table(rows=1, cols=len(cols))
    table.style = "Light Grid Accent 1"
    for cell, name in zip(table.rows[0].cells, cols):
        cell.text = name
        for run in cell.paragraphs[0].runs:
            run.bold = True

    for finding in findings:
        row = table.add_row().cells
        row[0].text = str(finding.get("location", "—"))
        row[1].text = str(finding.get("description", "—"))
        measured = finding.get("measured_thickness_mm")
        minimum = finding.get("minimum_thickness_mm")
        row[2].text = f"{measured} mm" if measured is not None else "—"
        row[3].text = f"{minimum} mm" if minimum is not None else "—"
        row[4].text = str(finding.get("severity", "—"))


def _calculations(doc: Document, calcs: list[dict[str, Any]]) -> None:
    """Render every step. This is the PS's "calculations with steps shown"."""
    if not calcs:
        return
    _add_heading(doc, "Calculations", level=2)
    for calc in calcs:
        label = calc.get("description") or calc.get("formula") or "Calculation"
        para = doc.add_paragraph()
        para.add_run(f"{label}: ").bold = True
        para.add_run(str(calc.get("formatted", calc.get("value", "—"))))

        for step in calc.get("steps", []):
            step_para = doc.add_paragraph(str(step.get("text", "")),
                                          style="List Bullet")
            for run in step_para.runs:
                run.font.size = Pt(9)


def _citations(doc: Document, citations: list[dict[str, Any]]) -> None:
    if not citations:
        return
    _add_heading(doc, "Sources", level=2)
    for index, cite in enumerate(citations, start=1):
        source = cite.get("document", "unknown document")
        page = cite.get("page")
        locator = f", p. {page}" if page is not None else ""
        para = doc.add_paragraph(f"[C{index}] {source}{locator}")
        for run in para.runs:
            run.font.size = Pt(9)


def docx_render(
    ctx: ToolContext,
    *,
    title: str,
    summary: str = "",
    sections: list[dict[str, Any]] | None = None,
    findings: list[dict[str, Any]] | None = None,
    calculations: list[dict[str, Any]] | None = None,
    citations: list[dict[str, Any]] | None = None,
    recommendation: str = "",
    filename: str | None = None,
) -> Observation:
    """Build a Word deliverable from a typed spec."""
    doc = Document()

    _add_heading(doc, title, level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Prepared {date.today().isoformat()} · ").italic = True
    meta.add_run("DRAFT — requires human review and approval").italic = True

    if summary:
        _add_heading(doc, "Summary", level=2)
        doc.add_paragraph(summary)

    if findings:
        _add_heading(doc, "Findings", level=2)
        _findings_table(doc, findings)

    for section in sections or []:
        _add_heading(doc, section.get("heading", "Section"), level=2)
        for para in section.get("paragraphs", []):
            doc.add_paragraph(str(para))

    _calculations(doc, calculations or [])

    if recommendation:
        _add_heading(doc, "Recommendation", level=2)
        doc.add_paragraph(recommendation)

    _citations(doc, citations or [])

    doc.add_paragraph()
    footer = doc.add_paragraph(
        "Generated on-premise by PRAHARI. No data left the organisation's "
        "perimeter in producing this document."
    )
    for run in footer.runs:
        run.font.size = Pt(8)
        run.italic = True

    safe = filename or f"{title.lower().replace(' ', '_')[:48]}_{uuid.uuid4().hex[:6]}.docx"
    if not safe.endswith(".docx"):
        safe += ".docx"

    out_dir = Path(ctx.workspace) / "artifacts"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / safe
    doc.save(out_path)

    return Observation(
        ok=True,
        summary=f"rendered {safe}",
        data={
            "artifact_kind": "docx",
            "filename": safe,
            "path": str(out_path),
            "relative_path": f"artifacts/{safe}",
            "title": title,
            "findings_count": len(findings or []),
            "calculations_count": len(calculations or []),
            "citations_count": len(citations or []),
        },
    )


def xlsx_render(
    ctx: ToolContext,
    *,
    title: str,
    sheets: list[dict[str, Any]],
    filename: str | None = None,
) -> Observation:
    """Build an Excel deliverable."""
    from openpyxl import Workbook
    wb = Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    for sheet_data in sheets:
        sheet_name = sheet_data.get("name", "Sheet")
        ws = wb.create_sheet(title=sheet_name[:31]) # Excel limits to 31 chars
        
        rows = sheet_data.get("rows", [])
        for r_idx, row in enumerate(rows, 1):
            if isinstance(row, list):
                for c_idx, val in enumerate(row, 1):
                    ws.cell(row=r_idx, column=c_idx, value=str(val) if val is not None else "")
            elif isinstance(row, dict):
                for c_idx, (k, val) in enumerate(row.items(), 1):
                    # Write header on first row if we're passing dicts
                    if r_idx == 1:
                        ws.cell(row=1, column=c_idx, value=str(k))
                        ws.cell(row=2, column=c_idx, value=str(val) if val is not None else "")
                    else:
                        ws.cell(row=r_idx+1, column=c_idx, value=str(val) if val is not None else "")
                        
    safe = filename or f"{title.lower().replace(' ', '_')[:48]}_{uuid.uuid4().hex[:6]}.xlsx"
    if not safe.endswith(".xlsx"):
        safe += ".xlsx"

    out_dir = Path(ctx.workspace) / "artifacts"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / safe
    wb.save(out_path)

    return Observation(
        ok=True,
        summary=f"rendered {safe}",
        data={
            "artifact_kind": "xlsx",
            "filename": safe,
            "path": str(out_path),
            "relative_path": f"artifacts/{safe}",
            "title": title,
        },
    )


def pptx_render(
    ctx: ToolContext,
    *,
    title: str,
    slides: list[dict[str, Any]],
    filename: str | None = None,
) -> Observation:
    """Build a PowerPoint deliverable."""
    from pptx import Presentation
    prs = Presentation()
    
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title_shape = slide.shapes.title
    subtitle = slide.placeholders[1]
    title_shape.text = title
    subtitle.text = f"Prepared {date.today().isoformat()}"
    
    # Content slides
    bullet_slide_layout = prs.slide_layouts[1]
    for slide_data in slides:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = slide_data.get("title", "Slide")
        
        tf = body_shape.text_frame
        points = slide_data.get("bullets", [])
        if points:
            tf.text = str(points[0])
            for point in points[1:]:
                p = tf.add_paragraph()
                p.text = str(point)

    safe = filename or f"{title.lower().replace(' ', '_')[:48]}_{uuid.uuid4().hex[:6]}.pptx"
    if not safe.endswith(".pptx"):
        safe += ".pptx"

    out_dir = Path(ctx.workspace) / "artifacts"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / safe
    prs.save(out_path)

    return Observation(
        ok=True,
        summary=f"rendered {safe}",
        data={
            "artifact_kind": "pptx",
            "filename": safe,
            "path": str(out_path),
            "relative_path": f"artifacts/{safe}",
            "title": title,
        },
    )


registry.register(
    ToolSpec(
        name="docx.render",
        description=(
            "Create a Word document deliverable (approval note, report) from "
            "structured content. Use this to produce the final file."
        ),
        args_schema={
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Document title"},
                "summary": {"type": "string", "description": "Opening summary paragraph"},
                "sections": {"type": "array",
                             "description": "[{heading, paragraphs[]}]"},
                "findings": {"type": "array", "description": "Finding objects"},
                "calculations": {"type": "array",
                                 "description": "Results from calc.evaluate"},
                "citations": {"type": "array", "description": "[{document, page}]"},
                "recommendation": {"type": "string",
                                   "description": "Closing recommendation"},
                "filename": {"type": "string", "description": "Optional output name"},
            },
            "required": ["title"],
        },
        capabilities=frozenset({CAP_RENDER}),
        handler=docx_render,
        returns="path to the generated .docx in the run workspace",
    )
)

registry.register(
    ToolSpec(
        name="xlsx.render",
        description="Create an Excel spreadsheet deliverable.",
        args_schema={
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Spreadsheet title"},
                "sheets": {
                    "type": "array",
                    "description": "List of sheets. Format: [{name, rows: [[col1, col2, ...]]}]"
                },
                "filename": {"type": "string", "description": "Optional output name"},
            },
            "required": ["title", "sheets"],
        },
        capabilities=frozenset({CAP_RENDER}),
        handler=xlsx_render,
        returns="path to the generated .xlsx in the run workspace",
    )
)

registry.register(
    ToolSpec(
        name="pptx.render",
        description="Create a PowerPoint presentation deliverable.",
        args_schema={
            "type": "object",
            "properties": {
                "title": {"type": "string", "description": "Presentation title"},
                "slides": {
                    "type": "array",
                    "description": "List of slides. Format: [{title, bullets: [...]}]"
                },
                "filename": {"type": "string", "description": "Optional output name"},
            },
            "required": ["title", "slides"],
        },
        capabilities=frozenset({CAP_RENDER}),
        handler=pptx_render,
        returns="path to the generated .pptx in the run workspace",
    )
)
